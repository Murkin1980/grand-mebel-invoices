import shutil
import openpyxl
import os

source = r'D:\ИП Гранд Мебель\Накладные 2026\Продажа № 1 от 20.01.2026.xlsx'
target = r'instance/template_waybill.xlsx'

if not os.path.exists('instance'):
    os.makedirs('instance')

shutil.copy2(source, target)

wb = openpyxl.load_workbook(target)
ws = wb.active

# В оригинале товары на строках 24, 25, 26.
# Удаляем строки 25, 26
ws.delete_rows(25, 2)

# Wipe all cells in row 24 that contain data
for col in range(1, 40):
    try:
        ws.cell(row=24, column=col).value = None
    except AttributeError:
        pass

# Wipe totals which are now at row 25
for col in range(1, 40):
    try:
        val = str(ws.cell(row=25, column=col).value or '')
        if '336900' in val or '4' in val:
            ws.cell(row=25, column=col).value = None
            
        val2 = str(ws.cell(row=26, column=col).value or '')
        if 'Четыре' in val2:
            ws.cell(row=26, column=col).value = None
    except AttributeError:
        pass

wb.save(target)
print("Template created at instance/template_waybill.xlsx")
