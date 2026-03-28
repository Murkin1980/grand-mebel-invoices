# -*- coding: utf-8 -*-
import os
import re
from flask import Flask, render_template, request, redirect, url_for, flash, send_file, jsonify
from flask_sqlalchemy import SQLAlchemy
from datetime import datetime, date
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

app = Flask(__name__)
app.config['SECRET_KEY'] = os.environ.get('SECRET_KEY', 'invoice-app-secret-key-2026')

# База данных: поддерживаем SQLite локально и PostgreSQL на Render
db_url = os.environ.get('DATABASE_URL', '')
if db_url.startswith('postgres://'):
    db_url = db_url.replace('postgres://', 'postgresql://', 1)
app.config['SQLALCHEMY_DATABASE_URI'] = db_url or 'sqlite:///invoice_app.db'
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['ESF_CERT_PATH'] = os.environ.get('ESF_CERT_PATH', '')
app.config['ESF_CERT_PASSWORD'] = os.environ.get('ESF_CERT_PASSWORD', 'Aa123456')
app.config['ESF_SENDER_BIN'] = os.environ.get('ESF_SENDER_BIN', '910226302322')
db = SQLAlchemy(app)


# ─── Models ───────────────────────────────────────────────────────────────────

class Contractor(db.Model):
    __tablename__ = 'contractors'
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(200), nullable=False)
    bin_iin = db.Column(db.String(50), nullable=False)
    address = db.Column(db.String(500), nullable=False)
    invoices = db.relationship('Invoice', backref='contractor', lazy=True)

    def __repr__(self):
        return f'{self.name} ({self.bin_iin})'


class Product(db.Model):
    __tablename__ = 'products'
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(300), nullable=False)
    unit = db.Column(db.String(50), default='Штука')
    default_price = db.Column(db.Float, default=0)
    is_service = db.Column(db.Boolean, default=False)

    def __repr__(self):
        return self.name


class InvoiceItem(db.Model):
    __tablename__ = 'invoice_items'
    id = db.Column(db.Integer, primary_key=True)
    invoice_id = db.Column(db.Integer, db.ForeignKey('invoices.id'), nullable=False)
    product_id = db.Column(db.Integer, db.ForeignKey('products.id'), nullable=False)
    quantity = db.Column(db.Float, default=1)
    price = db.Column(db.Float, nullable=False)
    product = db.relationship('Product')
    invoice = db.relationship('Invoice', backref='items')


class Invoice(db.Model):
    __tablename__ = 'invoices'
    id = db.Column(db.Integer, primary_key=True)
    number = db.Column(db.Integer, nullable=False)
    date = db.Column(db.Date, nullable=False)
    contractor_id = db.Column(db.Integer, db.ForeignKey('contractors.id'), nullable=False)
    contract_info = db.Column(db.String(300), default='Без договора')
    created_at = db.Column(db.DateTime, default=datetime.utcnow)

    def total(self):
        return sum(item.quantity * item.price for item in self.items)

    def total_text(self):
        total = int(self.total())
        return number_to_text(total)


# ─── Helpers ──────────────────────────────────────────────────────────────────

def number_to_text(num):
    """Конвертация числа в текст (тенге)"""
    if num == 0:
        return 'ноль'

    units_m = ['', 'один', 'два', 'три', 'четыре', 'пять', 'шесть', 'семь', 'восемь', 'девять']
    units_f = ['', 'одна', 'две', 'три', 'четыре', 'пять', 'шесть', 'семь', 'восемь', 'девять']
    teens = ['десять', 'одиннадцать', 'двенадцать', 'тринадцать', 'четырнадцать', 'пятнадцать',
             'шестнадцать', 'семнадцать', 'восемнадцать', 'девятнадцать']
    tens = ['', '', 'двадцать', 'тридцать', 'сорок', 'пятьдесят', 'шестьдесят', 'семьдесят', 'восемьдесят', 'девяносто']
    hundreds = ['', 'сто', 'двести', 'триста', 'четыреста', 'пятьсот', 'шестьсот', 'семьсот', 'восемьсот', 'девятьсот']

    def three_digits(n, feminine=False):
        result = []
        if n >= 100:
            result.append(hundreds[n // 100])
        r = n % 100
        if r >= 10 and r <= 19:
            result.append(teens[r - 10])
        else:
            if r >= 20:
                result.append(tens[r // 10])
            u = r % 10
            if u > 0:
                result.append(units_f[u] if feminine else units_m[u])
        return result

    result = []
    millions = num // 1_000_000
    thousands = (num % 1_000_000) // 1_000
    remainder = num % 1_000

    if millions:
        parts = three_digits(millions)
        result.extend(parts)
        if millions % 100 in range(11, 20):
            result.append('миллионов')
        elif millions % 10 == 1:
            result.append('миллион')
        elif millions % 10 in (2, 3, 4):
            result.append('миллиона')
        else:
            result.append('миллионов')

    if thousands:
        parts = three_digits(thousands, feminine=True)
        result.extend(parts)
        if thousands % 100 in range(11, 20):
            result.append('тысяч')
        elif thousands % 10 == 1:
            result.append('тысяча')
        elif thousands % 10 in (2, 3, 4):
            result.append('тысячи')
        else:
            result.append('тысяч')

    if remainder:
        result.extend(three_digits(remainder))

    return ' '.join(r for r in result if r).strip()


# ─── Routes: Index ─────────────────────────────────────────────────────────────

@app.route('/')
def index():
    search = request.args.get('search', '').strip()
    query = Invoice.query.order_by(Invoice.date.desc())
    if search:
        query = query.join(Contractor).filter(
            db.or_(
                Contractor.name.ilike(f'%{search}%'),
                db.cast(Invoice.number, db.String).ilike(f'%{search}%')
            )
        )
    invoices = query.all()
    return render_template('index.html', invoices=invoices, search=search)


# ─── Routes: Contractors ───────────────────────────────────────────────────────

@app.route('/contractors')
def contractors():
    search = request.args.get('search', '').strip()
    query = Contractor.query
    if search:
        query = query.filter(
            db.or_(
                Contractor.name.ilike(f'%{search}%'),
                Contractor.bin_iin.ilike(f'%{search}%')
            )
        )
    contractors_list = query.order_by(Contractor.name).all()
    return render_template('contractors.html', contractors=contractors_list, search=search)


@app.route('/contractors/add', methods=['POST'])
def add_contractor():
    name = request.form.get('name', '').strip()
    bin_iin = request.form.get('bin_iin', '').strip()
    address = request.form.get('address', '').strip()
    if not name or not bin_iin:
        flash('Название и БИН/ИИН обязательны', 'danger')
        return redirect(url_for('contractors'))
    contractor = Contractor(name=name, bin_iin=bin_iin, address=address)
    db.session.add(contractor)
    db.session.commit()
    flash(f'Контрагент «{name}» добавлен', 'success')
    return redirect(url_for('contractors'))


@app.route('/contractors/edit/<int:id>', methods=['GET', 'POST'])
def edit_contractor(id):
    contractor = Contractor.query.get_or_404(id)
    if request.method == 'POST':
        contractor.name = request.form.get('name', '').strip()
        contractor.bin_iin = request.form.get('bin_iin', '').strip()
        contractor.address = request.form.get('address', '').strip()
        db.session.commit()
        flash('Контрагент обновлён', 'success')
        return redirect(url_for('contractors'))
    return render_template('edit_contractor.html', contractor=contractor)


@app.route('/contractors/delete/<int:id>')
def delete_contractor(id):
    contractor = Contractor.query.get_or_404(id)
    if contractor.invoices:
        flash(f'Нельзя удалить — у контрагента есть счета ({len(contractor.invoices)} шт.)', 'danger')
        return redirect(url_for('contractors'))
    db.session.delete(contractor)
    db.session.commit()
    flash('Контрагент удалён', 'success')
    return redirect(url_for('contractors'))


# ─── Routes: Products ──────────────────────────────────────────────────────────

@app.route('/products')
def products():
    search = request.args.get('search', '').strip()
    query = Product.query
    if search:
        query = query.filter(Product.name.ilike(f'%{search}%'))
    products_list = query.order_by(Product.name).all()
    return render_template('products.html', products=products_list, search=search)


@app.route('/products/add', methods=['POST'])
def add_product():
    name = request.form.get('name', '').strip()
    unit = request.form.get('unit', 'Штука')
    default_price = float(request.form.get('default_price', 0) or 0)
    is_service = request.form.get('is_service') == 'on'
    if not name:
        flash('Введите название товара', 'danger')
        return redirect(url_for('products'))
    product = Product(name=name, unit=unit, default_price=default_price, is_service=is_service)
    db.session.add(product)
    db.session.commit()
    flash(f'«{name}» добавлен', 'success')
    return redirect(url_for('products'))


@app.route('/products/edit/<int:id>', methods=['GET', 'POST'])
def edit_product(id):
    product = Product.query.get_or_404(id)
    if request.method == 'POST':
        product.name = request.form.get('name', '').strip()
        product.unit = request.form.get('unit', 'Штука')
        product.default_price = float(request.form.get('default_price', 0) or 0)
        product.is_service = request.form.get('is_service') == 'on'
        db.session.commit()
        flash('Товар обновлён', 'success')
        return redirect(url_for('products'))
    return render_template('edit_product.html', product=product)


@app.route('/products/delete/<int:id>')
def delete_product(id):
    product = Product.query.get_or_404(id)
    db.session.delete(product)
    db.session.commit()
    flash('Товар удалён', 'success')
    return redirect(url_for('products'))


# ─── Routes: Invoices ──────────────────────────────────────────────────────────

@app.route('/invoice/new', methods=['GET', 'POST'])
def new_invoice():
    contractors_list = Contractor.query.order_by(Contractor.name).all()
    products_list = Product.query.order_by(Product.name).all()
    products_json = [{'id': p.id, 'name': p.name, 'unit': p.unit, 'default_price': p.default_price} for p in products_list]

    if request.method == 'POST':
        contractor_id = request.form.get('contractor_id')
        date_str = request.form.get('date')
        contract_info = request.form.get('contract_info', 'Без договора') or 'Без договора'

        last_invoice = Invoice.query.order_by(Invoice.number.desc()).first()
        next_number = (last_invoice.number + 1) if last_invoice else 1

        invoice = Invoice(
            number=next_number,
            date=datetime.strptime(date_str, '%Y-%m-%d').date(),
            contractor_id=contractor_id,
            contract_info=contract_info
        )
        db.session.add(invoice)
        db.session.flush()

        product_ids = request.form.getlist('product_id')
        quantities = request.form.getlist('quantity')
        prices = request.form.getlist('price')
        custom_names = request.form.getlist('custom_product_name')

        for i, product_id in enumerate(product_ids):
            qty = float(quantities[i]) if i < len(quantities) and quantities[i] else 1
            prc = float(prices[i]) if i < len(prices) and prices[i] else 0
            if product_id == '__custom__':
                custom_name = custom_names[i] if i < len(custom_names) and custom_names[i] else 'Товар'
                temp_product = Product(name=custom_name, unit='Штука', default_price=prc)
                db.session.add(temp_product)
                db.session.flush()
                item = InvoiceItem(invoice_id=invoice.id, product_id=temp_product.id, quantity=qty, price=prc)
            elif product_id:
                item = InvoiceItem(invoice_id=invoice.id, product_id=int(product_id), quantity=qty, price=prc)
            else:
                continue
            db.session.add(item)

        db.session.commit()
        flash(f'Счёт №{invoice.number} создан', 'success')
        return redirect(url_for('view_invoice', id=invoice.id))

    return render_template('new_invoice.html',
                           contractors=contractors_list,
                           products=products_list,
                           products_json=products_json,
                           today=date.today().isoformat())


@app.route('/invoice/<int:id>')
def view_invoice(id):
    invoice = Invoice.query.get_or_404(id)
    return render_template('invoice_view.html', invoice=invoice)


@app.route('/invoice/edit/<int:id>', methods=['GET', 'POST'])
def edit_invoice(id):
    invoice = Invoice.query.get_or_404(id)
    contractors_list = Contractor.query.order_by(Contractor.name).all()
    products_list = Product.query.order_by(Product.name).all()
    products_json = [{'id': p.id, 'name': p.name, 'unit': p.unit, 'default_price': p.default_price} for p in products_list]

    if request.method == 'POST':
        invoice.contractor_id = request.form.get('contractor_id')
        invoice.date = datetime.strptime(request.form.get('date'), '%Y-%m-%d').date()
        invoice.contract_info = request.form.get('contract_info', 'Без договора') or 'Без договора'

        # Delete old items
        InvoiceItem.query.filter_by(invoice_id=invoice.id).delete()

        product_ids = request.form.getlist('product_id')
        quantities = request.form.getlist('quantity')
        prices = request.form.getlist('price')
        custom_names = request.form.getlist('custom_product_name')

        for i, product_id in enumerate(product_ids):
            qty = float(quantities[i]) if i < len(quantities) and quantities[i] else 1
            prc = float(prices[i]) if i < len(prices) and prices[i] else 0
            if product_id == '__custom__':
                custom_name = custom_names[i] if i < len(custom_names) and custom_names[i] else 'Товар'
                temp_product = Product(name=custom_name, unit='Штука', default_price=prc)
                db.session.add(temp_product)
                db.session.flush()
                item = InvoiceItem(invoice_id=invoice.id, product_id=temp_product.id, quantity=qty, price=prc)
            elif product_id:
                item = InvoiceItem(invoice_id=invoice.id, product_id=int(product_id), quantity=qty, price=prc)
            else:
                continue
            db.session.add(item)

        db.session.commit()
        flash(f'Счёт №{invoice.number} обновлён', 'success')
        return redirect(url_for('view_invoice', id=invoice.id))

    return render_template('edit_invoice.html',
                           invoice=invoice,
                           contractors=contractors_list,
                           products=products_list,
                           products_json=products_json)


@app.route('/invoice/delete/<int:id>')
def delete_invoice(id):
    invoice = Invoice.query.get_or_404(id)
    num = invoice.number
    InvoiceItem.query.filter_by(invoice_id=id).delete()
    db.session.delete(invoice)
    db.session.commit()
    flash(f'Счёт №{num} удалён', 'success')
    return redirect(url_for('index'))


# ─── Routes: Export ────────────────────────────────────────────────────────────

@app.route('/invoice/export/<int:id>')
def export_invoice(id):
    invoice = Invoice.query.get_or_404(id)

    wb = Workbook()
    ws = wb.active
    ws.title = "Счет на оплату"

    thin = Side(style='thin')
    thin_border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
    total_fill = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
    bold = Font(bold=True)

    ws.page_setup.orientation = 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_margins.left = 0.5
    ws.page_margins.right = 0.5
    ws.page_margins.top = 0.75
    ws.page_margins.bottom = 0.75

    # Row 1 — notice
    ws.merge_cells('A1:G1')
    ws['A1'] = 'ВНИМАНИЕ! Оплата данного счета означает согласие с условиями поставки товара. Уведомление об оплате обязательно, в противном случае не гарантируется наличие товара на складе.'
    ws['A1'].font = Font(size=8, italic=True)
    ws['A1'].alignment = Alignment(wrap_text=True, horizontal='center')
    ws.row_dimensions[1].height = 28

    # Row 3 — title
    ws.merge_cells('A3:D3')
    ws['A3'] = 'СЧЕТ НА ОПЛАТУ'
    ws['A3'].font = Font(bold=True, size=16)
    ws['A3'].alignment = Alignment(horizontal='center')
    ws.merge_cells('E3:G3')
    ws['E3'] = 'Без печати и подписи не действителен'
    ws['E3'].font = Font(size=8, italic=True, color='FF0000')
    ws['E3'].alignment = Alignment(horizontal='right')

    # Rows 5-7 — реквизиты
    ws['A5'] = 'Бенефициар:'
    ws.merge_cells('B5:D5')
    ws['B5'] = 'ИП "ГРАНД МЕБЕЛЬ"'
    ws['B5'].font = bold
    ws['E5'] = 'ИИК:'
    ws.merge_cells('F5:G5')
    ws['F5'] = 'KZ84722S000035000586'

    ws['A6'] = 'БИН:'
    ws.merge_cells('B6:D6')
    ws['B6'] = '910226302322'
    ws['E6'] = 'Кбе:'
    ws['F6'] = '19'

    ws['A7'] = 'Банк:'
    ws.merge_cells('B7:D7')
    ws['B7'] = 'АО «Kaspi Bank»'
    ws['E7'] = 'БИК:'
    ws['F7'] = 'CASPKZKA'

    # Row 9 — номер и дата счёта
    date_str = invoice.date.strftime('%d.%m.%Y')
    ws.merge_cells('A9:D9')
    ws['A9'] = f'Счёт на оплату № {invoice.number}'
    ws['A9'].font = Font(bold=True, size=14)
    ws.merge_cells('E9:G9')
    ws['E9'] = f'от {date_str} г.'
    ws['E9'].font = Font(bold=True, size=12)

    # Rows 11-13 — поставщик, покупатель, договор
    ws['A11'] = 'Поставщик:'
    ws['A11'].font = bold
    ws.merge_cells('B11:G11')
    ws['B11'] = 'ИП "ГРАНД МЕБЕЛЬ" | ИИН 910226302322 | Казахстан, Аулиеагаш, МИКРОРАЙОН МАДЕНИЕТ, УЛИЦА ТАСБОЛАТ, дом 34'
    ws['B11'].font = Font(size=9)
    ws['B11'].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[11].height = 22

    ws['A12'] = 'Покупатель:'
    ws['A12'].font = bold
    ws.merge_cells('B12:G12')
    ws['B12'] = f'{invoice.contractor.name} | {invoice.contractor.bin_iin} | {invoice.contractor.address}'
    ws['B12'].font = Font(size=9)
    ws['B12'].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[12].height = 22

    ws['A13'] = 'Договор:'
    ws['A13'].font = bold
    ws.merge_cells('B13:G13')
    ws['B13'] = invoice.contract_info

    # Row 15 — table headers
    table_start = 15
    headers = ['№', 'Наименование товара (работ, услуг)', 'Кол-во', 'Ед. изм.', 'Цена (KZT)', 'Сумма (KZT)']
    col_widths = [5, 42, 9, 9, 14, 14]
    cols = ['A', 'B', 'C', 'D', 'E', 'F']

    for col, (header, width, letter) in enumerate(zip(headers, col_widths, cols), 1):
        cell = ws.cell(row=table_start, column=col)
        cell.value = header
        cell.font = Font(bold=True, size=10)
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[letter].width = width
    ws.row_dimensions[table_start].height = 30

    # Merge G column with F for headers
    ws.merge_cells(f'F{table_start}:G{table_start}')

    row = table_start + 1
    for idx, item in enumerate(invoice.items, 1):
        values = [idx, item.product.name, item.quantity, item.product.unit, item.price, item.quantity * item.price]
        aligns = ['center', 'left', 'center', 'center', 'right', 'right']
        for col, (val, align) in enumerate(zip(values, aligns), 1):
            cell = ws.cell(row=row, column=col)
            cell.value = val
            cell.border = thin_border
            cell.alignment = Alignment(horizontal=align, vertical='center')
            if col in (5, 6):
                cell.number_format = '#,##0.00'
        ws.merge_cells(f'F{row}:G{row}')
        ws.row_dimensions[row].height = 18
        row += 1

    # Total row
    total = invoice.total()
    ws.merge_cells(f'A{row}:E{row}')
    ws[f'A{row}'] = 'ИТОГО:'
    ws[f'A{row}'].font = Font(bold=True, size=11)
    ws[f'A{row}'].alignment = Alignment(horizontal='right')
    ws[f'A{row}'].fill = total_fill
    ws[f'A{row}'].border = thin_border
    ws.merge_cells(f'F{row}:G{row}')
    ws[f'F{row}'] = total
    ws[f'F{row}'].font = Font(bold=True, size=11)
    ws[f'F{row}'].fill = total_fill
    ws[f'F{row}'].border = thin_border
    ws[f'F{row}'].number_format = '#,##0.00'
    ws[f'F{row}'].alignment = Alignment(horizontal='right')
    ws.row_dimensions[row].height = 20
    row += 1

    ws[f'A{row}'] = 'В том числе НДС:'
    ws[f'A{row}'].font = Font(size=9)
    ws[f'C{row}'] = 'Без НДС'
    ws[f'C{row}'].font = Font(size=9)
    row += 1

    ws.merge_cells(f'A{row}:C{row}')
    ws[f'A{row}'] = f'Всего наименований {len(invoice.items)}, на сумму'
    ws[f'A{row}'].font = Font(size=10, bold=True)
    ws.merge_cells(f'D{row}:G{row}')
    ws[f'D{row}'] = f'{total:,.2f} KZT'.replace(',', ' ')
    ws[f'D{row}'].font = Font(size=10, bold=True)
    row += 1

    total_text = invoice.total_text()
    ws.merge_cells(f'A{row}:B{row}')
    ws[f'A{row}'] = 'Всего к оплате:'
    ws[f'A{row}'].font = Font(size=11, bold=True)
    ws.merge_cells(f'C{row}:G{row}')
    ws[f'C{row}'] = f'{total_text.title()} тенге 00 тиын'
    ws[f'C{row}'].font = Font(size=11, bold=True)
    row += 2

    ws.merge_cells(f'A{row}:C{row}')
    ws[f'A{row}'] = 'Условия оплаты:'
    ws[f'A{row}'].font = bold
    row += 1
    ws.merge_cells(f'A{row}:G{row}')
    ws[f'A{row}'] = 'Оплата в течение 5 банковских дней. Уведомление об оплате обязательно.'
    ws[f'A{row}'].font = Font(size=9)
    row += 2

    # Signature
    ws.merge_cells(f'A{row}:C{row}')
    ws[f'A{row}'] = 'ПОСТАВЩИК:'
    ws[f'A{row}'].font = bold
    ws.merge_cells(f'E{row}:G{row}')
    ws[f'E{row}'] = 'ПОКУПАТЕЛЬ:'
    ws[f'E{row}'].font = bold
    row += 1
    ws.merge_cells(f'A{row}:C{row}')
    ws[f'A{row}'] = 'ИП "ГРАНД МЕБЕЛЬ"'
    ws[f'A{row}'].font = bold
    ws.merge_cells(f'E{row}:G{row}')
    ws[f'E{row}'] = invoice.contractor.name
    ws[f'E{row}'].font = bold
    row += 3
    ws.merge_cells(f'A{row}:C{row}')
    ws[f'A{row}'] = '_______________________'
    ws[f'A{row}'].alignment = Alignment(horizontal='center')
    ws.merge_cells(f'E{row}:G{row}')
    ws[f'E{row}'] = '_______________________'
    ws[f'E{row}'].alignment = Alignment(horizontal='center')
    row += 1
    ws.merge_cells(f'A{row}:C{row}')
    ws[f'A{row}'] = 'М.П.  (подпись / Ф.И.О.)'
    ws[f'A{row}'].font = Font(size=8, italic=True)
    ws[f'A{row}'].alignment = Alignment(horizontal='center')
    ws.merge_cells(f'E{row}:G{row}')
    ws[f'E{row}'] = 'М.П.  (подпись / Ф.И.О.)'
    ws[f'E{row}'].font = Font(size=8, italic=True)
    ws[f'E{row}'].alignment = Alignment(horizontal='center')

    ws.column_dimensions['G'].width = 2

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f'Счет_№{invoice.number}_от_{date_str}.xlsx'
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename
    )


# ─── Routes: ESF ───────────────────────────────────────────────────────────────

@app.route('/esf/settings', methods=['GET', 'POST'])
def esf_settings():
    if request.method == 'POST':
        app.config['ESF_CERT_PATH'] = request.form.get('cert_path', '')
        app.config['ESF_CERT_PASSWORD'] = request.form.get('cert_password', '')
        app.config['ESF_SENDER_BIN'] = request.form.get('sender_bin', '910226302322')
        flash('Настройки ЭСФ сохранены', 'success')
        return redirect(url_for('index'))
    return render_template('esf_settings.html',
                           cert_path=app.config.get('ESF_CERT_PATH', ''),
                           sender_bin=app.config.get('ESF_SENDER_BIN', '910226302322'))


@app.route('/esf/export/<int:id>')
def export_esf_xml(id):
    from esf_client.esf_models import Invoice as ESFInvoice, InvoiceItem as ESFInvoiceItem, Participant
    invoice = Invoice.query.get_or_404(id)
    seller = Participant(bin_iin='910226302322', name='ИП "ГРАНД МЕБЕЛЬ"',
                         address='Казахстан, Аулиеагаш, МИКРОРАЙОН МАДЕНИЕТ, УЛИЦА ТАСБОЛАТ, дом 34',
                         bank_account='KZ84722S000035000586', bank_name='АО «Kaspi Bank»', bik='CASPKZKA')
    buyer = Participant(bin_iin=invoice.contractor.bin_iin, name=invoice.contractor.name, address=invoice.contractor.address)
    items = []
    for item in invoice.items:
        unit_code = '999' if item.product.unit.lower() in ['услуга', 'услуги'] else '166'
        items.append(ESFInvoiceItem(name=item.product.name, unit_code=unit_code,
                                    unit_name=item.product.unit, quantity=item.quantity,
                                    price=item.price, tax_rate=0))
    esf_invoice = ESFInvoice(invoice_number=f"ЭСФ-{invoice.number}", invoice_date=invoice.date,
                              seller=seller, buyer=buyer, shipment_date=invoice.date)
    esf_invoice.items = items
    xml_content = esf_invoice.to_xml_string()
    filename = f'ESF_{invoice.number}_{invoice.date.strftime("%d%m%Y")}.xml'
    return send_file(BytesIO(xml_content.encode('utf-8')), mimetype='application/xml',
                     as_attachment=True, download_name=filename)


@app.route('/esf/send/<int:id>')
def send_to_esf(id):
    from esf_client.esf_client import MockESFClient
    invoice = Invoice.query.get_or_404(id)
    cert_path = app.config.get('ESF_CERT_PATH', '')
    if cert_path and os.path.exists(cert_path):
        try:
            from esf_client.esf_client import ESFClient
            client = ESFClient(cert_path, app.config.get('ESF_CERT_PASSWORD', ''), app.config.get('ESF_SENDER_BIN', ''))
            client.create_session()
            from esf_client.esf_models import create_invoice_from_order
            esf_inv = create_invoice_from_order(invoice)
            result = client.send_invoice(esf_inv.to_xml_string())
            client.close_session()
            if result.get('success'):
                flash(f'ЭСФ отправлен! UUID: {result.get("uuid", "N/A")}', 'success')
            else:
                flash(f'Ошибка: {result.get("error")}', 'danger')
        except Exception as e:
            flash(f'Ошибка: {str(e)}', 'danger')
    else:
        client = MockESFClient()
        result = client.send_invoice('')
        flash(f'Тестовый режим: {result.get("message")}', 'info')
    return redirect(url_for('view_invoice', id=id))


# ─── API ───────────────────────────────────────────────────────────────────────

@app.route('/api/product/<int:id>')
def api_product(id):
    p = Product.query.get_or_404(id)
    return jsonify({'id': p.id, 'name': p.name, 'unit': p.unit, 'price': p.default_price})


# ─── DB Init ───────────────────────────────────────────────────────────────────

def init_db():
    with app.app_context():
        db.create_all()
        if Contractor.query.count() == 0:
            contractors_data = [
                ('ИП "Ай-Ханым"', '811102401040', 'г.Алматы, Наурызбайский р-н, мкр. Акжар, ул.Дәулеткерей, здание 1/18'),
                ('ТОО "QULPYNAI BAKERY"', '201140032253', 'г. Алматы, БЦ Кворум, мкр. Мамыр-1, дом № 26'),
                ('ТОО "InterHouse"', '180640029936', 'г. Алматы, ул. Армянская, д. 7А'),
                ('ИП «Qulpynai»', '810219300694', 'г. Алматы, микрорайон Мамыр-7, дом 21А, кв. 70'),
                ('TOO «Бута Груп»', '941240002645', 'Алматы, м-н Нуркент, дом 9, корпус 18, кв. 46'),
                ('ТОО «Рем Маш»', '141040003420', 'г. Алматы, Жетысуский р-н, ул. Жумабаева 232 В.'),
                ('ТОО "CAR ID"', '220840021610', 'г. Алматы, ул. Шашкина, д. 24'),
                ('ИП КазСтеклоПроф', '960601400677', 'Боралдай, УЛИЦА ЖҰМАДІЛОВ, дом 22А'),
            ]
            for name, bin_iin, address in contractors_data:
                db.session.add(Contractor(name=name, bin_iin=bin_iin, address=address))

        if Product.query.count() == 0:
            products_data = [
                ('Тумба под колонку', 'Штука', 26700), ('Шкаф раздевалка Grey 1800*1050*600', 'Штука', 93730),
                ('Шкаф раздевалка Grey Tandem 1800*1050*600', 'Штука', 122740), ('Шкафчик навесной 916*616', 'Штука', 27300),
                ('Стол металокаркас 1400х700х750', 'Штука', 80320), ('Тумба с 2мя ящиками на колесах', 'Штука', 57130),
                ('Тумба с 3-мя ящиками на колесах', 'Штука', 59460), ('Стол тумба с 3мя ящиками', 'Штука', 68500),
                ('Кровать 1000х2000 с мягкой панелью', 'Штука', 122500), ('Стол перегородка', 'Штука', 154000),
                ('Шкаф стеллаж', 'Штука', 206500), ('Шкаф лофт', 'Штука', 136500),
                ('Кухонный гарнитур Olive VDF', 'Штука', 505564), ('Шкаф 2700*1500*600', 'Штука', 700700),
                ('Шкафчик 1000х880', 'Штука', 71200), ('Реставрация шкафов', 'Услуга', 46500),
                ('Стеновая панель из ЛДСП 138 м²', 'Услуга', 2555000), ('Распил', 'Услуга', 30000),
                ('Шкаф архивный Blum', 'Штука', 289960), ('Стол 1500х900', 'Штука', 49320),
                ('Стол полка 1500х400', 'Штука', 57520), ('Тумба Olive VDF', 'Штука', 36500),
                ('Боковина Шалфей', 'Штука', 5250), ('Столешница искусственный камень добор', 'Штука', 22000),
            ]
            for name, unit, price in products_data:
                db.session.add(Product(name=name, unit=unit, default_price=price, is_service=(unit == 'Услуга')))
        db.session.commit()


if __name__ == '__main__':
    init_db()
    app.run(debug=True, port=5000)
