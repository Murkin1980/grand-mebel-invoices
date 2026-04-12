import os
import openpyxl

files = [f for f in os.listdir('.') if f.endswith('.xlsx')]
if files:
    print(f"Opening: {files[0]}")
    wb = openpyxl.load_workbook(files[0], data_only=True)
    ws = wb.active
    
    with open('excel_out.txt', 'w', encoding='utf-8') as f:
        for row in ws.iter_rows(min_row=1, max_row=30, values_only=True):
            f.write(str([str(cell)[:100] if cell is not None else '' for cell in row[:25]]) + '\n')
        
        f.write("--- Items Table ---\n")
        in_table = False
        for row in ws.iter_rows(min_row=1, max_row=50, values_only=True):
            if row[0] == '№':
                in_table = True
                continue
            if in_table:
                if not row[0] or str(row[0]).startswith('ИТОГО'):
                    break
                f.write(str([str(cell)[:100] if cell is not None else '' for cell in row[:25]]) + '\n')
else:
    print("No xlsx files found")
